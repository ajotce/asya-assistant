import io

from fastapi.testclient import TestClient

from app.api.deps_auth import get_db_session
from app.core.config import get_settings
from app.main import app
from app.storage.runtime import file_store
from tests.auth_helpers import override_db_session, setup_test_db


def _create_session(client: TestClient) -> str:
    response = client.post("/api/session")
    assert response.status_code == 201
    return response.json()["session_id"]


def _authed_client(tmp_path, monkeypatch) -> TestClient:
    _, engine = setup_test_db(tmp_path, monkeypatch)
    app.dependency_overrides[get_db_session] = override_db_session(engine)
    client = TestClient(app)
    client.post(
        "/api/auth/register",
        json={"email": "files@example.com", "display_name": "Files", "password": "strong-pass-123"},
    )
    client.post("/api/auth/login", json={"email": "files@example.com", "password": "strong-pass-123"})
    return client


def test_upload_rejects_more_than_max_files_per_message(tmp_path, monkeypatch) -> None:
    client = _authed_client(tmp_path, monkeypatch)
    session_id = _create_session(client)

    files = [("files", (f"f{i}.pdf", b"%PDF-1.7", "application/pdf")) for i in range(11)]
    response = client.post(f"/api/session/{session_id}/files", files=files)

    assert response.status_code == 400
    assert "не более 10 файлов" in response.json()["detail"]
    app.dependency_overrides.clear()


def test_upload_rejects_unsupported_file_type(tmp_path, monkeypatch) -> None:
    client = _authed_client(tmp_path, monkeypatch)
    session_id = _create_session(client)

    response = client.post(
        f"/api/session/{session_id}/files",
        files=[("files", ("notes.txt", b"hello", "text/plain"))],
    )

    assert response.status_code == 400
    assert "не поддерживается" in response.json()["detail"]
    app.dependency_overrides.clear()


def test_upload_rejects_file_over_size_limit(tmp_path, monkeypatch) -> None:
    client = _authed_client(tmp_path, monkeypatch)
    session_id = _create_session(client)

    settings = get_settings()
    old_limit = settings.max_file_size_mb
    settings.max_file_size_mb = 1
    try:
        oversized = b"x" * (1024 * 1024 + 1)
        response = client.post(
            f"/api/session/{session_id}/files",
            files=[("files", ("big.pdf", oversized, "application/pdf"))],
        )
    finally:
        settings.max_file_size_mb = old_limit

    assert response.status_code == 400
    assert "превышает лимит 1 МБ" in response.json()["detail"]
    app.dependency_overrides.clear()


def test_extract_text_from_pdf_docx_xlsx(tmp_path, monkeypatch) -> None:
    client = _authed_client(tmp_path, monkeypatch)
    session_id = _create_session(client)
    settings = get_settings()
    old_key = settings.vsellm_api_key
    settings.vsellm_api_key = "test-key"

    files = [
        ("files", ("sample.pdf", _make_pdf_bytes("PDF hello"), "application/pdf")),
        (
            "files",
            (
                "sample.docx",
                _make_docx_bytes("DOCX привет"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        ),
        (
            "files",
            (
                "sample.xlsx",
                _make_xlsx_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        ),
    ]
    try:
        from app.services import vsellm_client as vsellm_client_module

        old_post = vsellm_client_module.httpx.post

        def fake_post(*args, **kwargs):
            inputs = kwargs["json"]["input"]
            data = [{"embedding": [float(i + 1), 0.0, 0.0]} for i in range(len(inputs))]
            import httpx

            request = httpx.Request("POST", "https://api.vsellm.ru/v1/embeddings")
            return httpx.Response(status_code=200, request=request, json={"data": data})

        vsellm_client_module.httpx.post = fake_post
        response = client.post(f"/api/session/{session_id}/files", files=files)
    finally:
        vsellm_client_module.httpx.post = old_post
        settings.vsellm_api_key = old_key

    assert response.status_code == 201

    stored = file_store.get_session_files(session_id)
    by_name = {file.filename: file for file in stored}
    assert "PDF hello" in (by_name["sample.pdf"].extracted_text or "")
    assert "DOCX привет" in (by_name["sample.docx"].extracted_text or "")
    xlsx_text = by_name["sample.xlsx"].extracted_text or ""
    assert "[Лист: Sheet1]" in xlsx_text
    assert "Строка 1:" in xlsx_text
    assert "1=Имя" in xlsx_text
    app.dependency_overrides.clear()


def test_upload_rejects_damaged_docx(tmp_path, monkeypatch) -> None:
    client = _authed_client(tmp_path, monkeypatch)
    session_id = _create_session(client)

    response = client.post(
        f"/api/session/{session_id}/files",
        files=[("files", ("broken.docx", b"not-a-docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert response.status_code == 400
    assert "может быть поврежд" in response.json()["detail"]
    app.dependency_overrides.clear()


def test_upload_rejects_empty_xlsx(tmp_path, monkeypatch) -> None:
    client = _authed_client(tmp_path, monkeypatch)
    session_id = _create_session(client)

    response = client.post(
        f"/api/session/{session_id}/files",
        files=[("files", ("empty.xlsx", _make_empty_xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))],
    )
    assert response.status_code == 400
    assert "не содержит данных" in response.json()["detail"]
    app.dependency_overrides.clear()


def test_upload_returns_clear_error_when_embeddings_api_unavailable(tmp_path, monkeypatch) -> None:
    client = _authed_client(tmp_path, monkeypatch)
    session_id = _create_session(client)
    settings = get_settings()
    old_key = settings.vsellm_api_key
    settings.vsellm_api_key = "test-key"
    try:
        from app.services import vsellm_client as vsellm_client_module

        old_post = vsellm_client_module.httpx.post

        def fake_post(*args, **kwargs):
            import httpx

            request = httpx.Request("POST", "https://api.vsellm.ru/v1/embeddings")
            return httpx.Response(status_code=503, request=request, json={"error": "down"})

        vsellm_client_module.httpx.post = fake_post
        response = client.post(
            f"/api/session/{session_id}/files",
            files=[("files", ("sample.docx", _make_docx_bytes("text"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
        )
    finally:
        vsellm_client_module.httpx.post = old_post
        settings.vsellm_api_key = old_key

    assert response.status_code == 502
    assert "embeddings" in response.json()["detail"].lower()
    app.dependency_overrides.clear()


def test_upload_file_isolation_between_users(tmp_path, monkeypatch) -> None:
    _, engine = setup_test_db(tmp_path, monkeypatch)
    app.dependency_overrides[get_db_session] = override_db_session(engine)
    client_a = TestClient(app)
    client_b = TestClient(app)

    client_a.post(
        "/api/auth/register",
        json={"email": "fa@example.com", "display_name": "FA", "password": "strong-pass-123"},
    )
    client_a.post("/api/auth/login", json={"email": "fa@example.com", "password": "strong-pass-123"})
    session_id = _create_session(client_a)

    client_b.post(
        "/api/auth/register",
        json={"email": "fb@example.com", "display_name": "FB", "password": "strong-pass-123"},
    )
    client_b.post("/api/auth/login", json={"email": "fb@example.com", "password": "strong-pass-123"})
    response = client_b.post(
        f"/api/session/{session_id}/files",
        files=[("files", ("notes.txt", b"hello", "text/plain"))],
    )
    assert response.status_code == 404
    app.dependency_overrides.clear()


def _make_pdf_bytes(text: str) -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _make_docx_bytes(text: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _make_xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Имя", "Значение"])
    ws.append(["A", 123])
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


def _make_empty_xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()
